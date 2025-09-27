import os
from langgraph.graph import StateGraph, END
from typing import TypedDict, Annotated, Sequence, List, Dict
from langchain_core.messages import BaseMessage, SystemMessage, HumanMessage, ToolMessage
from operator import add as add_messages
from langchain_openai import ChatOpenAI
from langchain_core.tools import tool
from langchain_core.runnables import RunnableConfig

from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from pathlib import Path

from dotenv import load_dotenv

import pandas as pd
import sqlite3

load_dotenv()

DB_NAME = os.getenv("DB_NAME")
MODEL_OPENAI = os.getenv("MODEL_OPENAI")
API_KEY_OPENAI = os.getenv("API_KEY_OPENAI")
SYSTEM_PROMPT = os.getenv("SYSTEM_PROMPT")

conn = sqlite3.connect(DB_NAME, check_same_thread=False)

def save_with_chart(df):
    filepath = Path(__file__).parent / "test.xlsx"
    filepath = filepath.as_posix()

    wb = Workbook()
    sheet = wb.active

    for r in dataframe_to_rows(df, index=False, header=True):
        sheet.append(r)

    numeric_cols = [col for col in df.columns if pd.api.types.is_numeric_dtype(df[col])]
    if not numeric_cols:
        print("⚠️ Grafik uchun raqamli ustun topilmadi.")
        wb.save(filepath)
        return filepath

    first_num_col = df.columns.get_loc(numeric_cols[0]) + 1
    last_num_col = df.columns.get_loc(numeric_cols[-1]) + 1

    bar_chart = BarChart()
    bar_chart.varyColors = True
    data = Reference(sheet, min_row=1, max_row=sheet.max_row,
                     min_col=first_num_col, max_col=last_num_col)
    cats = Reference(sheet, min_col=1, min_row=2, max_row=sheet.max_row)

    bar_chart.add_data(data, titles_from_data=True)
    bar_chart.set_categories(cats)
    bar_chart.dataLabels = DataLabelList()
    bar_chart.dataLabels.showVal = True

    line_chart = LineChart()
    line_chart.add_data(data, titles_from_data=True)
    line_chart.set_categories(cats)

    bar_chart += line_chart

    last_col = sheet.max_column + 3
    sheet.add_chart(bar_chart, f"{get_column_letter(last_col)}2")

    wb.save(filepath)
    return filepath

@tool
def assistant_excel(prompt):
  """This function creates an Excel file depending on prompt."""
  results = pd.read_sql(prompt, conn)
  filename = save_with_chart(results)

  return {
        "reply": f"Excel tayyor",
        "_file": {
            "path": filename,
            "mime": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "caption": "file"
        }
    }

tools = [assistant_excel]

llm = ChatOpenAI(model=MODEL_OPENAI, temperature = 0, api_key=API_KEY_OPENAI).bind_tools(tools)

def concat_list(prev, new):
    return (prev or []) + (new or [])

class AgentState(TypedDict):
    messages: Annotated[Sequence[BaseMessage], add_messages]
    outbox: Annotated[List[Dict], concat_list]

def should_continue(state: AgentState):
    """Check if the last message contains tool calls."""
    result = state['messages'][-1]
    return hasattr(result, 'tool_calls') and len(result.tool_calls) > 0

tools_dict = {our_tool.name: our_tool for our_tool in tools}

def call_llm(state: AgentState, config: RunnableConfig):
    messages = [SystemMessage(content=SYSTEM_PROMPT)] + state['messages']

    response = llm.invoke(messages)

    return {"messages": [response]}

def take_action(state: AgentState, config: RunnableConfig) -> AgentState:
    """Execute tool calls from the LLM's response."""

    tool_calls = state['messages'][-1].tool_calls

    out_files: List[Dict] = []

    results = []
    for t in tool_calls:
        name = t['name']
        args = t['args']
        print(f"Calling Tool: {t['name']} with query: {t['args']}")

        if not t['name'] in tools_dict:
            print(f"\nTool: {t['name']} does not exist.")
            result = "Incorrect Tool Name, Please Retry and Select tool from List of Available tools."

        else:
            result = tools_dict[name].invoke({**args})

            reply_text = result.get("reply") if isinstance(result, dict) else str(result)

            if isinstance(result, dict) and "_file" in result:
                out_files.append(result["_file"])

            state['messages'].append(ToolMessage(tool_call_id=t['id'], name=t['name'], content=str(reply_text)))

    print("Tools Execution Complete. Back to the model!")
    if out_files:
        state["outbox"] = out_files

    return state

graph = StateGraph(AgentState)
graph.add_node("llm", call_llm)
graph.add_node("retriever_agent", take_action)

graph.add_conditional_edges(
    "llm",
    should_continue,
    {True: "retriever_agent", False: END}
)
graph.add_edge("retriever_agent", "llm")
graph.set_entry_point("llm")

rag_agent = graph.compile()